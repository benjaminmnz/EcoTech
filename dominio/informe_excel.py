from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from dominio.informe import Informe


class InformeExcel(Informe):
    def generar(self):
        self._preparar_destino(".xlsx")

        libro = Workbook()
        libro.remove(libro.active)

        resumen = libro.create_sheet("Resumen")
        resumen.append(["Informe de EcoTech Solutions"])
        resumen.append(["ID", self.get_id()])
        resumen.append(["Fecha", self.fecha_generacion.isoformat()])
        resumen.column_dimensions["A"].width = 32
        resumen.column_dimensions["B"].width = 20

        for titulo, encabezados, filas in self._obtener_secciones():
            hoja = libro.create_sheet(titulo)

            for numero_fila, fila in enumerate([encabezados] + filas, start=1):
                for numero_columna, valor in enumerate(fila, start=1):
                    celda = hoja.cell(row=numero_fila, column=numero_columna, value=valor)

                    if isinstance(valor, str):
                        celda.data_type = "s"

                    celda.alignment = Alignment(vertical="top", wrap_text=True)

                    if numero_fila == 1:
                        celda.font = Font(bold=True, color="FFFFFF")
                        celda.fill = PatternFill("solid", fgColor="28614D")
                    elif encabezados[numero_columna - 1] in ("Salario", "Horas", "Horas totales"):
                        celda.number_format = "#,##0.00"

            for numero_columna, encabezado in enumerate(encabezados, start=1):
                letra = get_column_letter(numero_columna)
                hoja.column_dimensions[letra].width = 12 if encabezado == "ID" else 30

            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions

        libro.save(str(self.ruta_archivo))
