"""Verificaciones reproducibles de ES2; no conservan los datos de prueba."""
from contextlib import contextmanager, redirect_stdout
from datetime import date
from decimal import Decimal
from io import StringIO
from pathlib import Path
from secrets import token_urlsafe
from tempfile import TemporaryDirectory
from unittest import TestCase, TestLoader, TestSuite, TextTestRunner
from unittest.mock import patch
from uuid import uuid4
import sys

import psycopg
from openpyxl import load_workbook

from dominio.persona import Persona
from dominio.empleado import Empleado
from dominio.departamento import Departamento
from dominio.proyecto import Proyecto
from dominio.registro_tiempo import RegistroTiempo
from dominio.cuenta_usuario import CuentaUsuario
from dominio.excepciones import ErrorPersistencia, AutorizacionError
from dominio.informe import Informe
from dominio.informe_pdf import InformePDF
from dominio.informe_excel import InformeExcel
from repositorio.conexion import obtener_conexion
from repositorio.departamento_repo import DepartamentoRepositorio
from repositorio.empleado_repo import EmpleadoRepositorio
from repositorio.proyecto_repo import ProyectoRepositorio
from repositorio.cuenta_usuario_repo import CuentaUsuarioRepositorio
from repositorio.registro_tiempo_repo import RegistroTiempoRepositorio
from repositorio.sistema_repo import SistemaRepositorio
from servicios import GestionEcoTech, Sesion, autenticar


def empleado_prueba(nombre="Empleado de prueba"):
    return Empleado(nombre, "Dirección de prueba", "912345678",
                    "prueba@example.com", 0, date(2026, 1, 1), 800000)


class DominioTest(TestCase):
    def test_herencia_y_salario(self):
        empleado = empleado_prueba()
        self.assertIsInstance(empleado, Persona)
        self.assertTrue(empleado.subir_salario(Decimal("0.10")))
        self.assertEqual(empleado.get_salario(), Decimal("800000.10"))
        for valor in (-1, 0, True, float("nan"), float("inf"), "texto"):
            self.assertFalse(empleado.subir_salario(valor))
        self.assertEqual(empleado.get_salario(), Decimal("800000.10"))

    def test_constructor_y_contacto_invalidos(self):
        for valor in (-1, 0, True, float("nan"), float("inf"), "texto"):
            with self.assertRaises(ValueError):
                Empleado("Nombre", "Dirección", "123", "a@example.com", 0, date.today(), valor)
        with self.assertRaises(ValueError):
            Persona(" ", "Dirección", "123", "a@example.com")
        with self.assertRaises(ValueError):
            Persona("Nombre", "Dirección", "123", "correo inválido")
        empleado = empleado_prueba()
        with self.assertRaises(ValueError):
            empleado.actualizar_datos_contacto("Otra dirección", " ")
        self.assertEqual(empleado.direccion, "Dirección de prueba")

    def test_horas_invalidas_no_modifican(self):
        registro = RegistroTiempo(1, date.today(), 4.5, "Tarea")
        for valor in (-1, 0, 25, True, False, float("nan"), float("inf"), "texto", 1.001):
            with self.assertRaises(ValueError):
                registro.registrar(valor)
        self.assertEqual(registro.get_horas_trabajadas(), Decimal("4.5"))

    def test_relaciones_y_coordinacion(self):
        empleado = empleado_prueba()
        otro = empleado_prueba("Otro")
        departamento = Departamento(1, "Departamento")
        departamento.asignar_empleado(empleado)
        with self.assertRaises(ValueError):
            Departamento(2, "Otro").asignar_empleado(empleado)
        copia = departamento.listar_empleados()
        copia.clear()
        self.assertEqual(len(departamento.listar_empleados()), 1)
        departamento.asignar_gerente(empleado)
        with self.assertRaises(ValueError):
            departamento.asignar_gerente(otro)
        proyecto = Proyecto(1, "Proyecto", "Descripción", date.today())
        with self.assertRaises(ValueError):
            proyecto.registrar_tiempo(empleado, 1, date.today(), 2, "Tarea")
        self.assertEqual(empleado.listar_registros(), [])
        proyecto.asignar_empleado(empleado)
        registro = proyecto.registrar_tiempo(empleado, 1, date.today(), 4.5, "Tarea")
        self.assertIs(registro.empleado, empleado)
        self.assertIs(registro.proyecto, proyecto)
        with self.assertRaises(ValueError):
            proyecto.registrar_tiempo(empleado, 1, date.today(), 2, "Duplicado")
        self.assertEqual(len(empleado.listar_registros()), 1)
        self.assertEqual(proyecto.calcular_horas_totales(), Decimal("4.5"))

    def test_hash_roles_y_propietario(self):
        clave = token_urlsafe(20)
        hash1 = CuentaUsuario.generar_hash(clave)
        hash2 = CuentaUsuario.generar_hash(clave)
        self.assertNotEqual(hash1, hash2)
        self.assertNotIn(clave, hash1)
        cuenta = CuentaUsuario(1, "usuario", hash1, "empleado")
        self.assertTrue(cuenta.verificar_contrasena(clave))
        self.assertFalse(cuenta.verificar_contrasena(clave + "x"))
        self.assertTrue(cuenta.autorizar("registro_tiempo"))
        self.assertFalse(cuenta.autorizar("empleados"))
        self.assertFalse(cuenta.autorizar("desconocido"))
        empleado_prueba().asignar_cuenta(cuenta)
        with self.assertRaises(ValueError):
            empleado_prueba("Otro").asignar_cuenta(cuenta)

    def test_informes_reales_y_polimorfismo(self):
        empleado = empleado_prueba()
        departamento = Departamento(1, "Departamento")
        departamento.asignar_empleado(empleado)
        proyecto = Proyecto(1, "Proyecto", "Descripción", date.today())
        proyecto.asignar_empleado(empleado)
        registro = proyecto.registrar_tiempo(empleado, 1, date.today(), 4.5, "Tarea <A> & B")
        with TemporaryDirectory() as carpeta:
            destino = Path(carpeta)
            informes = [
                InformePDF(1, [empleado], [proyecto], [registro], destino / "prueba.pdf"),
                InformeExcel(2, [empleado], [proyecto], [registro], destino / "prueba.xlsx")
            ]
            for informe in informes:
                informe.generar()
                with self.assertRaises(FileExistsError):
                    informe.generar()
            pdf = (destino / "prueba.pdf").read_bytes()
            self.assertTrue(pdf.startswith(b"%PDF"))
            self.assertIn(b"%%EOF", pdf[-1024:])
            libro = load_workbook(destino / "prueba.xlsx")
            self.assertEqual(libro["Proyectos"]["D2"].value, 4.5)
            self.assertEqual(libro["Registros"]["D2"].value, "Proyecto")
            self.assertEqual(libro["Registros"]["E2"].value, 4.5)
            libro.close()
            empleado.nombre = "=1+1"
            InformeExcel(3, [empleado], [], [], destino / "texto.xlsx").generar()
            libro = load_workbook(destino / "texto.xlsx")
            self.assertEqual(libro["Empleados"]["B2"].data_type, "s")
            libro.close()
            InformePDF(4, [], [], [], destino / "vacio.pdf").generar()
            InformeExcel(5, [], [], [], destino / "vacio.xlsx").generar()
        with self.assertRaises(TypeError):
            Informe(1, [], [], [], "base.pdf")

    def test_entrada_invalida_y_menu(self):
        import main
        with patch("builtins.input", side_effect=["texto", "0", "3"]), redirect_stdout(StringIO()):
            self.assertEqual(main.leer_id("ID"), 3)
        with patch("builtins.input", side_effect=["nan", "-1", "2,5"]), redirect_stdout(StringIO()):
            self.assertEqual(main.leer_decimal("Número"), Decimal("2.5"))
        with patch("builtins.input", side_effect=["2026-02-30", "2026-02-28"]), redirect_stdout(StringIO()):
            self.assertEqual(main.leer_fecha("Fecha"), date(2026, 2, 28))


class IntegracionTest(TestCase):
    def test_crud_seguridad_relaciones_rollback_y_permisos(self):
        marca = uuid4().hex[:10]
        clave = token_urlsafe(20)
        with obtener_conexion() as conexion:
            SistemaRepositorio(conexion).comprobar_esquema()
            # Todo este bloque se revierte incluso si las comprobaciones pasan.
            with conexion.transaction(force_rollback=True):
                with conexion.cursor() as cursor:
                    cursor.execute("SET CONSTRAINTS ALL IMMEDIATE")
                departamentos = DepartamentoRepositorio(conexion)
                empleados = EmpleadoRepositorio(conexion)
                proyectos = ProyectoRepositorio(conexion)
                cuentas = CuentaUsuarioRepositorio(conexion)
                registros = RegistroTiempoRepositorio(conexion)

                departamento = departamentos.crear("Verificación " + marca)
                self.assertEqual(departamentos.obtener_por_id(departamento.get_id()).nombre, departamento.nombre)
                departamento = departamentos.actualizar(departamento.get_id(), "Verificación actualizada " + marca)
                self.assertIn(departamento.get_id(), [d.get_id() for d in departamentos.listar()])
                administrador = empleados.crear(empleado_prueba("Admin " + marca), departamento.get_id())
                empleado = empleados.crear(empleado_prueba("Empleado " + marca), departamento.get_id())
                administrador.actualizar_datos_contacto("Contacto actualizado", "987654321")
                administrador = empleados.actualizar(administrador, departamento.get_id())
                self.assertEqual(administrador.direccion, "Contacto actualizado")
                departamentos.asignar_gerente(departamento.get_id(), administrador.get_id())

                hash_clave = CuentaUsuario.generar_hash(clave)
                ca = cuentas.crear(CuentaUsuario(0, "admin_" + marca, hash_clave, "administrador"), administrador.get_id())
                ce = cuentas.crear(CuentaUsuario(0, "empleado_" + marca, hash_clave, "empleado"), empleado.get_id())
                proyecto = proyectos.crear(Proyecto(0, "Proyecto " + marca, "Descripción", date.today()))
                proyecto.nombre = "Proyecto actualizado " + marca
                proyecto = proyectos.actualizar(proyecto)
                proyectos.asignar_empleado(proyecto.get_id(), administrador.get_id())
                proyectos.asignar_empleado(proyecto.get_id(), empleado.get_id())
                self.assertIn(empleado.get_id(), proyectos.listar_ids_empleados(proyecto.get_id()))
                registro = registros.crear(administrador.get_id(), proyecto.get_id(), date.today(), 4.5, "Tarea")
                registro = registros.actualizar(registro.get_id(), administrador.get_id(), proyecto.get_id(), date.today(), 3.5, "Corregida")
                self.assertEqual(registro.get_horas_trabajadas(), Decimal("3.5"))

                # El texto se busca exactamente; no convierte WHERE en una condición verdadera.
                payload = "' OR '1'='1"
                cantidad_antes = len(departamentos.listar())
                resultados = departamentos.buscar_por_nombre(payload)
                self.assertTrue(all(d.nombre == payload for d in resultados))
                self.assertNotIn(departamento.get_id(), [d.get_id() for d in resultados])
                literal = departamentos.crear(payload)
                resultados = departamentos.buscar_por_nombre(payload)
                self.assertIn(literal.get_id(), [d.get_id() for d in resultados])
                self.assertTrue(all(d.nombre == payload for d in resultados))
                departamentos.eliminar(literal.get_id())
                self.assertEqual(len(departamentos.listar()), cantidad_antes)

                # Fallo real en segunda escritura: nombre de usuario duplicado.
                cantidad_empleados = len(empleados.listar())
                with self.assertRaises(ErrorPersistencia):
                    with conexion.transaction():
                        nuevo = empleados.crear(empleado_prueba("Debe revertirse"), departamento.get_id())
                        cuentas.crear(CuentaUsuario(0, ca.usuario, hash_clave, "empleado"), nuevo.get_id())
                self.assertIsNone(empleados.obtener_por_id(nuevo.get_id()))
                self.assertEqual(len(empleados.listar()), cantidad_empleados)

                # La BD también rechaza horas si el empleado no está asignado.
                sin_asignar = proyectos.crear(Proyecto(0, "Sin asignación", "Prueba", date.today()))
                with self.assertRaises(ErrorPersistencia):
                    with conexion.transaction():
                        registros.crear(empleado.get_id(), sin_asignar.get_id(), date.today(), 1, "Inválida")

                # Conexión prestada con savepoints para probar servicios sin conservar escrituras.
                @contextmanager
                def prestada():
                    with conexion.transaction():
                        yield conexion

                with patch("servicios.obtener_conexion", prestada):
                    self.assertIsNone(autenticar(ca.usuario, clave + "incorrecta"))
                    self.assertEqual(autenticar(ca.usuario, clave).empleado_id, administrador.get_id())
                    ga = GestionEcoTech(Sesion(ca.get_id(), administrador.get_id()))
                    ge = GestionEcoTech(Sesion(ce.get_id(), empleado.get_id()))
                    with self.assertRaises(AutorizacionError):
                        ge.crear_departamento("No autorizado")
                    with self.assertRaises(AutorizacionError):
                        ge.datos_informes()
                    with self.assertRaises(AutorizacionError):
                        ge.crear_registro(administrador.get_id(), proyecto.get_id(), date.today(), 1, "Ajeno")
                    with self.assertRaises(AutorizacionError):
                        ge.actualizar_registro(registro.get_id(), date.today(), 2, "Ajeno")
                    propio = ge.crear_registro(empleado.get_id(), proyecto.get_id(), date.today(), 2, "Propio")
                    self.assertEqual([r.get_id() for r in ge.registros()], [propio.get_id()])
                    datos = ga.datos_informes()
                    proyecto_cargado = next(p for p in datos[2] if p.get_id() == proyecto.get_id())
                    self.assertEqual(proyecto_cargado.calcular_horas_totales(), Decimal("5.5"))
                    registro_cargado = next(r for r in datos[3] if r.get_id() == registro.get_id())
                    empleado_cargado = next(e for e in datos[1] if e.get_id() == administrador.get_id())
                    self.assertIs(registro_cargado.empleado, empleado_cargado)
                    self.assertIs(registro_cargado.proyecto, proyecto_cargado)
                    self.assertIs(empleado_cargado.departamento.get_gerente(), empleado_cargado)
                    self.assertEqual(empleado_cargado.get_cuenta().get_id(), ca.get_id())

                    # Crear empleado y cuenta por el servicio: el duplicado revierte al empleado.
                    cantidad = len(empleados.listar())
                    with self.assertRaises(ErrorPersistencia):
                        ga.crear_empleado(empleado_prueba("No guardar"), departamento.get_id(), ca.usuario, clave, "empleado")
                    self.assertEqual(len(empleados.listar()), cantidad)

                    # Prueba de consola autenticada y de opción no autorizada.
                    import main
                    salida = StringIO()
                    with patch("main.obtener_conexion", prestada), patch("main.necesita_administrador", return_value=False), patch("main.getpass", return_value=clave), patch("builtins.input", side_effect=[ce.usuario, "1", "3", "4", "1", "0", "0"]), patch.object(sys, "argv", ["main.py"]), redirect_stdout(salida):
                        self.assertEqual(main.main(), 0)
                    self.assertIn("Opción inválida o no autorizada", salida.getvalue())
                    self.assertIn("Propio", salida.getvalue())

                # No se permite quitar asignaciones con historia, ni borrar proyecto con horas.
                with self.assertRaises(ErrorPersistencia):
                    with conexion.transaction():
                        proyectos.quitar_empleado(proyecto.get_id(), empleado.get_id())
                with self.assertRaises(ErrorPersistencia):
                    with conexion.transaction():
                        proyectos.eliminar(proyecto.get_id())

                self.assertTrue(registros.eliminar(propio.get_id()))
                self.assertIsNone(registros.obtener_por_id(propio.get_id()))
                # Eliminar al gerente borra su cuenta/horas y quita su designación.
                self.assertTrue(empleados.eliminar(administrador.get_id()))
                self.assertIsNone(cuentas.obtener_por_empleado(administrador.get_id()))
                self.assertIsNone(registros.obtener_por_id(registro.get_id()))
                self.assertIsNone(dict(departamentos.listar_gerentes())[departamento.get_id()])
                proyectos.quitar_empleado(proyecto.get_id(), empleado.get_id())
                self.assertTrue(proyectos.eliminar(proyecto.get_id()))
                self.assertIsNone(proyectos.obtener_por_id(proyecto.get_id()))
                empleados.eliminar(empleado.get_id())
                self.assertTrue(departamentos.eliminar(departamento.get_id()))
                self.assertIsNone(departamentos.obtener_por_id(departamento.get_id()))

    def test_persistencia_entre_conexiones(self):
        identificador = None
        nombre = "Verificación reapertura " + uuid4().hex[:10]
        try:
            with obtener_conexion() as conexion:
                departamento = DepartamentoRepositorio(conexion).crear(nombre)
                identificador = departamento.get_id()
            with obtener_conexion() as conexion:
                recuperado = DepartamentoRepositorio(conexion).obtener_por_id(identificador)
                self.assertIsNotNone(recuperado)
                self.assertEqual(recuperado.nombre, nombre)
        finally:
            if identificador is not None:
                with obtener_conexion() as conexion:
                    DepartamentoRepositorio(conexion).eliminar(identificador)


def ejecutar_verificaciones(solo_local=False):
    suite = TestSuite()
    suite.addTests(TestLoader().loadTestsFromTestCase(DominioTest))
    if not solo_local:
        suite.addTests(TestLoader().loadTestsFromTestCase(IntegracionTest))
    print("Verificaciones EcoTech: dominio, exportaciones" + ("" if solo_local else ", SQL, rollback y permisos"))
    resultado = TextTestRunner(verbosity=2).run(suite)
    return 0 if resultado.wasSuccessful() else 1
